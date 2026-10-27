#!/usr/bin/env python3
"""Combine batch PPTX parts into one deck, keeping charts and pictures alive.

Copying a shape's XML alone leaves its relationship ids dangling (charts live
in their own part, pictures in `ppt/media/`), so every r:id in the copied XML is
resolved against the SOURCE part, attached to the DESTINATION part and written
back (audit of 2026-08-23). The slide background, which `add_slide()` resets to
white, is set again too. See `references/generation-workflow.md`.

Usage:
    python3 .claude/skills/pptx-generator/scripts/combine_decks.py \\
        --parts "$DECK_DIR/mydeck-part*.pptx" \\
        --out "$DECK_DIR/mydeck-final.pptx" \\
        --background 0E1116 \\
        --delete-parts

No test guards the combine; check a combined deck by eye after changing this file.
"""
from __future__ import annotations

import argparse
import copy
import glob
import re
import shutil
import sys
from pathlib import Path

RELS_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
BLANK_LAYOUT = 6

# The LAST run of digits in the stem. `(?!.*\d)` rejects any match with another
# digit after it, so `2026-08-31-deck-part12` reads 12 and not 2026.
_TRAILING_NUMBER_RE = re.compile(r"(\d+)(?!.*\d)")

# python-pptx is an optional extra (`x-heading-requires: ["documents"]`), so it
# is imported inside the functions that need it and importing this module stays
# free of it. The test that enforced this left with fc2a9c95.


def hex_to_rgb(hex_color: str):
    from pptx.dml.color import RGBColor

    h = hex_color.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def remap_relationships(element, src_part, dst_part) -> int:
    """Re-point every r:id in a copied shape at the destination part.

    Returns the number of relationships remapped, so a caller can assert the
    copy actually carried something rather than trusting it silently.
    """
    remapped = 0
    for node in element.iter():
        for name, rid in list(node.attrib.items()):
            if not name.startswith(f"{{{RELS_NS}}}"):
                continue
            rel = src_part.rels[rid]
            if rel.is_external:
                new_rid = dst_part.relate_to(
                    rel.target_ref, rel.reltype, is_external=True)
            else:
                new_rid = dst_part.relate_to(rel.target_part, rel.reltype)
            node.set(name, new_rid)
            remapped += 1
    return remapped


def copy_slide(src_slide, dst_prs, background: str | None) -> int:
    """Append one slide's shapes to a new slide in dst_prs."""
    new_slide = dst_prs.slides.add_slide(dst_prs.slide_layouts[BLANK_LAYOUT])

    # add_slide() gives a default white background, and shape copying does not
    # carry the slide's own background - it is a slide property, not a shape.
    if background:
        new_slide.background.fill.solid()
        new_slide.background.fill.fore_color.rgb = hex_to_rgb(background)

    remapped = 0
    for shape in src_slide.shapes:
        element = copy.deepcopy(shape.element)
        remapped += remap_relationships(
            element, src_slide.part, new_slide.part)
        new_slide.shapes._spTree.append(element)
    return remapped


def sort_part_files(paths: list[Path]) -> list[Path]:
    """Order part files by the number in their name, not lexically.

    Part names are unpadded (``{name}-part1.pptx``), so `sorted()` would put
    part10 before part2. A name with no digits is kept and sorts after every
    numbered part, among themselves by name.
    """
    def key(path: Path) -> tuple[int, int, str]:
        match = _TRAILING_NUMBER_RE.search(path.stem)
        if match is None:
            return (1, 0, path.name)
        return (0, int(match.group(1)), path.name)

    return sorted(paths, key=key)


def combine(part_files: list[Path], out_path: Path,
            background: str | None = None) -> Path:
    """Merge parts in the given order into out_path. Returns out_path."""
    from pptx import Presentation

    if not part_files:
        raise SystemExit("Error: no part files matched.")
    if len(part_files) == 1:
        shutil.copy(str(part_files[0]), str(out_path))
        return out_path

    combined = Presentation(str(part_files[0]))
    for part_file in part_files[1:]:
        part_prs = Presentation(str(part_file))
        for slide in part_prs.slides:
            copy_slide(slide, combined, background)
    combined.save(str(out_path))
    return out_path


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Combine batch PPTX parts into one deck.")
    parser.add_argument("--parts", required=True,
                        help="glob for the part files, e.g. '$DECK_DIR/deck-part*.pptx'")
    parser.add_argument("--out", required=True, help="path of the combined deck")
    parser.add_argument("--background",
                        help="brand background hex, e.g. 0E1116. Slides added by "
                             "the combine default to white without it.")
    parser.add_argument("--delete-parts", action="store_true",
                        help="remove the part files after a successful combine")
    args = parser.parse_args()

    part_files = sort_part_files([Path(p) for p in glob.glob(args.parts)])
    if not part_files:
        print(f"Error: no files matched {args.parts}", file=sys.stderr)
        return 2

    out_path = Path(args.out)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    combine(part_files, out_path, args.background)

    if args.delete_parts:
        for part_file in part_files:
            if part_file.resolve() != out_path.resolve():
                part_file.unlink()

    print(f"Combined {len(part_files)} part(s) -> {out_path}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
