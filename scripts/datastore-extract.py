#!/usr/bin/env python3
"""
DataStore Extraction Script for HEADING OS

Converts documents (PPTX, XLSX, PDF, DOCX, HTML) anywhere in the data
repository to readable markdown companion files (-extract.md). The companion is
the only surface the memory index can reach, so a document type this script
does not handle is a document the workspace cannot recall at all.

Scope since HOS-149 (2026-09-26), widened by C3 (2026-10-13) to every top-level
directory but the runtime store (`document_roots`): every directory under the roots
of the data repository, not only the seven DataStore homes, minus anything that
reads as an archive (see `_is_archive_name`), dot-directories, test fixture
folders (`_FIXTURE_DIRS`, since HOS-194), and the shared exclusions of
`scripts/utils/excluded_dirs.py`. A companion is never written
into an archive, never written as `<stem>.md`, and never written over a file
that is not itself a companion of the same source.

Each companion records the SHA-256 of its source and the converter that made
it, so a re-run regenerates only what is missing or stale. Stale means the
source's bytes changed, the converter NAME changed (the per-extension table
`CONVERTERS` is where a better converter swaps in), or the companion predates
the provenance line. A converter VERSION change is recorded and does not by
itself make a companion stale: a patch release of a parser would otherwise
re-embed the whole corpus.

Usage:
    python scripts/datastore-extract.py                   # every document dir of the data repo
    python scripts/datastore-extract.py <dir-or-file>     # one folder, or one document
    python scripts/datastore-extract.py --dry-run         # counts only, writes nothing
    python scripts/datastore-extract.py --update-index    # also update INDEX.md
    python scripts/datastore-extract.py --force           # regenerate every companion

Prerequisites:
    liteparse (core), markitdown, openpyxl, python-pptx, python-docx (the
    `documents` extra). All pinned in `pyproject.toml`.

    Which converter reads which format is the `CONVERTERS` table, and the
    choice is measured, not habitual. HOS-148, MEASURED 2026-09-26 on 22 of the
    operator's own documents, four converters each, judged blind: PDF goes
    through liteparse (best tables and completeness, the only one that reads a
    page with no text layer), DOCX, PPTX and HTML through markitdown (headings,
    lists, slide tables and notes kept; the only HTML reader of the four), with
    this file's own DOCX and PPTX readers as the fallback when markitdown
    raises, and XLSX stays on openpyxl, the only one of the four that lost
    nothing, hidden sheets included. The measurements and what was NOT tested
    are in the bake-off report in the data repository (registry name
    `system-work`, folder `2026-09-26_converter-bakeoff`).

    OCR runs only for a PDF page that has no text layer, with liteparse's local
    Tesseract in Russian and English (`scripts/utils/ocr.py` carries the
    measurement behind that choice). The document never leaves the machine.
    liteparse fetches a language MODEL from GitHub the first time it needs one
    it has not cached. An OCR failure is written into the companion.

Tests: none since the 2026-09-21 test-suite split. Two guarantees were held here
and both now rest on review: a rollback never deletes a file it did not first
back up, and the extractor reaches every format in the corpus rather than the
subset whose libraries happen to be importable. The second is lost the
moment a converter's package leaves `pyproject.toml` while its row stays in
`CONVERTERS`: the import fails per file, and the run names each failure.
"""

import argparse
import functools
import hashlib
import os
import re
import sys
import textwrap
from collections import Counter
from datetime import datetime
from importlib.metadata import version as dist_version
from pathlib import Path
from typing import Callable, NamedTuple

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from scripts.utils.atomic import atomic_write_text
from scripts.utils.colors import GREEN, YELLOW, RED, BOLD, RESET
from scripts.utils.data_path import data_path
from scripts.utils.excluded_dirs import NAMES as SHARED_EXCLUDED_NAMES
from scripts.utils.runtime_dir import store_dirname
from scripts.utils.data_homes import homes
from scripts.utils.ocr import OCR_LANGUAGE
from scripts.utils.workspace import get_default_tz, get_data_root


def datastore_homes() -> tuple[Path, ...]:
    """Every directory the DataStore's records live in, existing ones only."""
    return homes("datastore-homes")


def index_file():
    """INDEX.md of the DataStore, by its registry name `datastore-index`."""
    return data_path("datastore-index")


def extract_xlsx(filepath):
    """Extract XLSX content to markdown."""
    try:
        from openpyxl import load_workbook
    except ImportError:
        print(f"{RED}Error: openpyxl not installed. Run: uv sync --extra documents{RESET}")
        return None

    wb = load_workbook(filepath, data_only=True)
    lines = []
    lines.append(f"# Extract: {filepath.name}")
    lines.append(f"")
    lines.append(f"> Auto-extracted from `{filepath.name}` on {datetime.now(get_default_tz()).strftime('%Y-%m-%d')}")
    lines.append(f"> This is a companion file for Claude to read. The original XLSX is the source of truth.")
    lines.append(f"")

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        lines.append(f"## Sheet: {sheet_name}")
        lines.append(f"")

        rows = _xlsx_grid(ws.iter_rows(values_only=True))
        if not rows:
            lines.append("*(empty sheet)*")
            lines.append("")
            continue
        lines.extend(_markdown_table(rows))
        lines.append("")

    return "\n".join(lines)


def _xlsx_grid(rows) -> list:
    """The sheet's rows with every empty row and every empty column removed.

    openpyxl reports a sheet's full used range, and a formatted but blank cell
    extends that range, so a price list or compliance matrix came out as tens
    of thousands of `|  |  |` cells around a few hundred values. MEASURED
    2026-09-30 over the 214 XLSX companions outside archive directories in the
    data repository: 92% of their 60.7 million characters were spaces, pipes
    and dashes, and one
    companion was 31.3 MB of which 99.8% was empty cells. Only a column or row
    with no value anywhere is removed, so no value moves out of its column.
    Cells are no longer cut at 100 characters, as `_markdown_table` stopped
    doing for DOCX tables, and a row wider than the first one is no longer
    clipped to its width.
    """
    grid = [["" if c is None else str(c) for c in row] for row in rows]  # `_markdown_table` escapes
    grid = [row for row in grid if any(c.strip() for c in row)]
    if not grid:
        return []
    width = max(len(row) for row in grid)
    grid = [row + [""] * (width - len(row)) for row in grid]
    keep = [i for i in range(width) if any(row[i].strip() for row in grid)]
    return [[row[i] for i in keep] for row in grid]


def _cell(value) -> str:
    """One spreadsheet cell, safe inside a markdown table row.

    Cells were joined straight into `" | "`, so any value containing a pipe
    added a phantom column and any value containing a newline ended the row
    early — and the companion `-extract.md` stopped representing the sheet it
    claims to be an extract of.
    """
    if value is None:
        return ""
    text = str(value)
    return text.replace("\\", "\\\\").replace("|", "\\|").replace("\r", " ").replace("\n", " ")


@functools.lru_cache(maxsize=1)
def _markitdown():
    """One converter per process: importing markitdown costs about a second
    (MEASURED 2026-09-26, three runs, 1.08-1.14 s), and a datastore run
    converts hundreds of files."""
    from markitdown import MarkItDown
    return MarkItDown(enable_plugins=False)


def _markitdown_text(filepath, suffix: str) -> str:
    """markitdown's text for one file, read AS the format its suffix names.

    `MarkItDown.convert` guesses the format from the bytes and falls back to
    plain text, so a corrupt `.docx` came back as its own garbage and was
    written as a companion. MEASURED 2026-09-26: a 9-byte file that is not a
    zip produced a companion and exit 0, where the python-docx reader had
    raised and the run exited 1. A binary format therefore goes straight to
    its own converter, which raises on a file it cannot open. HTML is text,
    so it keeps the guessing path, which is also the one that detects a
    non-UTF-8 charset; the suffix is passed as a hint.
    """
    from markitdown import StreamInfo
    from markitdown.converters import DocxConverter, PptxConverter

    _markitdown()  # the one-time import cost, paid once per process
    info = StreamInfo(extension=suffix, filename=filepath.name, local_path=str(filepath))
    binary = {".docx": DocxConverter, ".pptx": PptxConverter}.get(suffix)
    if binary is None:
        return _markitdown().convert(str(filepath), stream_info=info).text_content
    with open(filepath, "rb") as stream:
        return binary().convert(stream, info).text_content


_EMPTY_TABLE_ROW = re.compile(r"^\|(\s*\|)+\s*$")
_TABLE_SEPARATOR = re.compile(r"^\|(\s*:?-{3,}:?\s*\|)+\s*$")


def _promote_empty_table_headers(text: str) -> str:
    """Lift a table's first data row into a header row that came out empty.

    markitdown writes every Word table whose first row is not marked as a
    header with an EMPTY header row, so each table opens with a line of blank
    cells. MEASURED 2026-09-26 in the HOS-148 bake-off: the blind judge marked
    it down on every DOCX. Only a table that starts with an all-blank row, then
    a separator, then a row is touched.
    """
    lines = text.split("\n")
    out = []
    i = 0
    while i < len(lines):
        starts_table = i == 0 or not lines[i - 1].lstrip().startswith("|")
        if (starts_table and i + 2 < len(lines)
                and _EMPTY_TABLE_ROW.match(lines[i])
                and _TABLE_SEPARATOR.match(lines[i + 1])
                and lines[i + 2].lstrip().startswith("|")):
            out.extend((lines[i + 2], lines[i + 1]))
            i += 3
            continue
        out.append(lines[i])
        i += 1
    return "\n".join(out)


def _markdown_companion(filepath, kind: str, text: str) -> str:
    """A companion around the Markdown markitdown produced, capped at a LINE.

    A line boundary, not a paragraph one. Table rows are separated by single
    newlines, so cutting at the last blank line threw away every row of a page
    that is one big table: MEASURED 2026-09-26 in the HOS-149 review, a
    heading, a paragraph and a 4,000-row table gave a 535-byte companion
    holding only the heading and the paragraph.

    Which converter made it is not written here: `_stamp` records it on the
    provenance line, which is also what staleness reads.
    """
    text = _promote_empty_table_headers(text or "").strip()
    if not text:
        return _no_text_companion(filepath, kind,
                                  "The converter returned no text at all.")

    note = None
    if len(text) > MAX_EXTRACT_CHARS:
        cut = text.rfind("\n", 0, MAX_EXTRACT_CHARS)
        cut = cut if cut > 0 else MAX_EXTRACT_CHARS
        lines_kept = text.count("\n", 0, cut) + 1
        total_lines = text.count("\n") + 1
        text = text[:cut].rstrip()
        note = _truncation_note(MAX_EXTRACT_CHARS,
                                f"line {lines_kept} of {total_lines}", kind)

    lines = _extract_header(filepath, kind)
    if note:
        lines += [note, ""]
    lines += [text, ""]
    if note:
        lines += [note, ""]
    return "\n".join(lines)


def extract_pptx(filepath):
    """PPTX to markdown through markitdown, this file's reader as the fallback.

    markitdown keeps slide boundaries, slide tables and speaker notes; it was
    the blind judge's pick on 2 of 3 decks in HOS-148. On the third it raised
    `ValueError` inside python-pptx over a shape position written as a decimal,
    so its failure is caught HERE, named, and the native reader runs instead.
    """
    try:
        text = _markitdown_text(filepath, ".pptx")
    except Exception as exc:  # noqa: BLE001 - named below, then the native reader runs
        reason = f"{type(exc).__name__}: {' '.join(str(exc).split())}"
        print(f"  {YELLOW}markitdown failed on {filepath.name} ({reason}); "
              f"using the native PPTX reader{RESET}")
        return _extract_pptx_native(filepath, reason)
    return _markdown_companion(filepath, "PPTX", text)


def _pptx_shape_lines(shape) -> list:
    """One shape as markdown lines: groups opened, tables kept as tables.

    Reading only `has_text_frame` shapes lost every slide table and every
    shape inside a group. MEASURED 2026-09-26 (HOS-148): all 9 tables of an
    organisation chart deck, all 4 of a partner framework deck, both of a
    product deck.
    """
    from pptx.shapes.group import GroupShape

    if isinstance(shape, GroupShape):
        return [line for inner in shape.shapes for line in _pptx_shape_lines(inner)]
    if getattr(shape, "has_table", False):
        rows = _markdown_table([[c.text for c in row.cells] for row in shape.table.rows])
        return ["", *rows, ""] if rows else []
    if getattr(shape, "has_text_frame", False):
        return [f"- {p.text.strip()}" for p in shape.text_frame.paragraphs if p.text.strip()]
    return []


def _extract_pptx_native(filepath, fallback_reason: str | None = None):
    """Extract PPTX text, tables and speaker notes with python-pptx."""
    try:
        from pptx import Presentation
    except ImportError:
        print(f"{RED}Error: python-pptx not installed. Run: uv sync --extra documents{RESET}")
        return None

    prs = Presentation(str(filepath))
    lines = _extract_header(filepath, "PPTX")
    if fallback_reason:
        # The provenance line names the table's converter, markitdown, because
        # that is what staleness compares; this line says what actually read it.
        lines += [f"> Read with python-pptx: markitdown failed ({fallback_reason}).", ""]
    lines.append(f"Total slides: {len(prs.slides)}")
    lines.append("")

    for i, slide in enumerate(prs.slides, 1):
        lines.append(f"## Slide {i}")
        lines.append("")

        texts = [line for shape in slide.shapes for line in _pptx_shape_lines(shape)]
        if any(t.strip() for t in texts):
            lines.extend(texts)
        else:
            lines.append("*(no text content - may be an image/diagram slide)*")

        frame = slide.notes_slide.notes_text_frame if slide.has_notes_slide else None
        if frame is not None:
            notes = frame.text.strip()
            if notes:
                lines.append("")
                lines.append("### Notes")
                lines.append("")
                lines.append(notes)

        lines.append("")

    return "\n".join(lines)


# The ceiling on one companion's extracted text, in characters.
#
# 100,000 characters is roughly 15,000 words, or about 25,000 tokens: larger
# than any hand-written note in this workspace, and still small enough that an
# agent can read the whole companion in one go. The number exists because the
# companion feeds a shared memory index. A single 900-page manual extracted
# whole would out-mass every other document in the datastore put together, and
# every unrelated query would start retrieving chunks of it. Truncation is the
# lesser harm as long as it is VISIBLE, which is why the marker is written at
# both ends of the companion rather than only where the cut happened.
MAX_EXTRACT_CHARS = 100_000

# Pages per OCR call. Small enough that OCR stops close to the cap, large enough
# that LiteParse's per-call start-up is not paid on every page.
OCR_BATCH_PAGES = 8



def _extract_header(filepath, kind: str) -> list:
    """The header every companion carries, whatever the source format was."""
    return [
        f"# Extract: {filepath.name}",
        "",
        f"> Auto-extracted from `{filepath.name}` on "
        f"{datetime.now(get_default_tz()).strftime('%Y-%m-%d')}",
        f"> This is a companion file for Claude to read. The original {kind} "
        f"is the source of truth.",
        "",
    ]


def _no_text_companion(filepath, kind: str, detail: str) -> str:
    """A companion for a binary that yielded no text at all.

    The alternative (write nothing) was rejected. A companion holding only the
    header is worse than useless: it reads to the next agent as "this document
    is empty", which is a false statement about a file that may be the most
    important one in its folder. Writing no file at all is quieter but leaves
    the same document unreachable forever, with every later run re-opening the
    same scan to re-discover the same nothing.

    So the companion is written and SAYS SO. It is indexable, it names the
    binary as the thing to open. Since HOS-148 a PDF page without text is
    OCRed first, so this companion now means OCR also found nothing, or failed,
    and the detail says which.
    """
    lines = _extract_header(filepath, kind)
    lines.append(f"**No extractable text layer.** {detail}")
    lines.append("")
    lines.append(
        f"This {kind} carries no text this extractor can read. Open the "
        f"original `{filepath.name}` directly; it is the source of truth. This "
        f"companion is NOT the record of an empty document."
    )
    lines.append("")
    return "\n".join(lines)


def _truncation_note(cap: int, where: str, kind: str) -> str:
    return (f"> TRUNCATED at the {cap}-character cap, after {where}. This "
            f"companion is PARTIAL; open the original {kind} for the rest.")


def extract_pdf(filepath):
    """Extract PDF text to markdown, one section per page.

    A PDF with no text layer gets a companion that says so rather than an empty
    one or no file at all; the reasoning is in `_no_text_companion`.

    The extract is capped at `MAX_EXTRACT_CHARS` and cut at a PAGE boundary,
    never mid-page, so no section of the companion is a half sentence. The
    cap's own reasoning sits on the constant.

    LiteParse reads the text layer first with OCR OFF, and OCRs only the pages
    that came back empty. OCR on every page is its default and it is what made
    it slow: MEASURED 2026-09-26 on a 57-page proposal, 60.5 s with OCR against
    0.3 s without, the difference being text inside diagrams.

    The parse itself is not caught here. An encrypted or corrupt PDF raises out
    of `parse`, and `scan_report` already names and counts a per-file parser
    failure without letting it end the batch. Catching it here would report the
    file as extracted-but-empty, which is the one outcome that must not happen.
    OCR is different: it is caught, and what went wrong is written into the
    companion.
    """
    try:
        from liteparse import LiteParse
    except ImportError:
        print(f"{RED}Error: liteparse not installed. Run: uv sync{RESET}")
        return None

    pages = LiteParse(ocr_enabled=False, quiet=True).parse(str(filepath)).pages
    texts = {page.page_num: _tidy_layout(page.text) for page in pages}
    pages_read = len(texts)
    blank = sorted(n for n, t in texts.items() if not t)
    ocr_pages: set = set()
    ocr_attempted: set = set()
    ocr_error = None

    def ocr_from(page_number: int) -> None:
        """OCR the next batch of blank pages, starting at `page_number`.

        In batches, in page order, and only while the companion still has room:
        the cap below cuts at a page anyway, so OCR past it is paid for and
        thrown away. MEASURED 2026-09-26 on a 206-page Russian book scan, OCR
        of every blank page up front took 5 min 25 s and 1.66 GB of memory,
        for a companion the cap cut after page 46.
        """
        nonlocal ocr_error
        batch = [n for n in blank if n >= page_number][:OCR_BATCH_PAGES]
        ocr_attempted.update(batch)
        try:
            scanned = LiteParse(
                ocr_enabled=True, ocr_failure_fatal=True, ocr_language=OCR_LANGUAGE,
                quiet=True,
                target_pages=",".join(str(n) for n in batch),
            ).parse(str(filepath)).pages
        except Exception as exc:  # noqa: BLE001 - written into the companion below
            ocr_error = f"{type(exc).__name__}: {exc}"
            print(f"  {YELLOW}OCR failed on {filepath.name}: {ocr_error}{RESET}")
            return
        for page in scanned:
            text = _tidy_layout(page.text)
            if text and page.page_num in texts and not texts[page.page_num]:
                texts[page.page_num] = text
                ocr_pages.add(page.page_num)

    body: list = []
    total_chars = 0
    truncated_after = None
    for page_number in sorted(texts):
        if total_chars >= MAX_EXTRACT_CHARS:
            truncated_after = page_number - 1
            break
        if (not texts[page_number] and page_number not in ocr_attempted
                and ocr_error is None):
            ocr_from(page_number)
        text = texts[page_number]
        if not text:
            # An image-only page OCR could not read either.
            continue
        body.append(f"## Page {page_number}"
                    + (" (OCR)" if page_number in ocr_pages else ""))
        body.append("")
        body.append(text)
        body.append("")
        total_chars += len(text)

    if not body:
        detail = (f"{pages_read} page(s) were read and none of them carried "
                  f"text, which usually means a scan or an image-only document.")
        if ocr_error:
            detail += f" OCR was attempted and failed: {ocr_error}"
        else:
            detail += " OCR was attempted and found no text."
        return _no_text_companion(filepath, "PDF", detail)

    note = None
    if truncated_after is not None:
        note = _truncation_note(MAX_EXTRACT_CHARS,
                                f"page {truncated_after}", "PDF")

    lines = _extract_header(filepath, "PDF")
    if note:
        lines.append(note)
        lines.append("")
    lines.append(f"Pages read: {pages_read}")
    if ocr_pages:
        lines.append(f"Pages read by OCR: {len(ocr_pages)}")
    if ocr_error:
        lines.append(f"OCR failed; {len(blank) - len(ocr_pages)} page(s) "
                     f"without a text layer have no text here: {ocr_error}")
    lines.append("")
    lines.extend(body)
    if note:
        lines.append(note)
        lines.append("")

    return "\n".join(lines)


def _tidy_layout(text: str) -> str:
    """LiteParse page text, kept in its layout and made safe for Markdown.

    LiteParse places text by its position on the page, which is what keeps a
    table's columns readable, and it indents every line of a page the same.
    That shared indent is removed (four spaces would turn the page into a code
    block for any Markdown reader), trailing spaces go, and runs of blank lines
    shrink to one.
    """
    lines = [line.rstrip() for line in (text or "").splitlines()]
    text = textwrap.dedent("\n".join(lines)).strip("\n")
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def _markdown_table(rows) -> list:
    """Rows of cell text as markdown table lines, the first row as the header.

    Cells are no longer cut at 100 characters: MEASURED 2026-09-26 the blind
    judge in HOS-148 counted 10 and 30 cut cells in two DOCX tender documents,
    and the cut was the text a reader came for.
    """
    lines: list = []
    for row in rows:
        cells = [_cell(c) for c in row]
        if not cells:
            continue
        lines.append("| " + " | ".join(cells) + " |")
        if len(lines) == 1:
            lines.append("| " + " | ".join(["---"] * len(cells)) + " |")
    return lines


def _docx_table_rows(table) -> list:
    """One DOCX table as markdown rows, the first row read as the header."""
    return _markdown_table([c.text for c in row.cells] for row in table.rows)


def extract_docx(filepath):
    """DOCX to markdown through markitdown, this file's reader as the fallback.

    markitdown keeps headings and lists, which the native reader below does
    not: MEASURED 2026-09-26 in HOS-148 it scored 1 of 5 on structure on every
    DOCX, with no heading at all in a document holding 190 of them. A markitdown
    failure is caught HERE, named, and the native reader runs instead.
    """
    try:
        text = _markitdown_text(filepath, ".docx")
    except Exception as exc:  # noqa: BLE001 - named below, then the native reader runs
        reason = f"{type(exc).__name__}: {' '.join(str(exc).split())}"
        print(f"  {YELLOW}markitdown failed on {filepath.name} ({reason}); "
              f"using the native DOCX reader{RESET}")
        return _extract_docx_native(filepath, reason)
    return _markdown_companion(filepath, "DOCX", text)


def extract_markitdown(filepath):
    """HTML to markdown through markitdown, the only reader of the four
    measured in HOS-148 that accepts HTML at all.

    markitdown drops `<script>` and `<style>`, renders tables as markdown
    tables, and shortens an inline `data:` URI to its scheme, so a report
    carrying base64 images does not become megabytes of noise. MEASURED
    2026-09-26 on markitdown 0.1.8 with a one-pixel PNG inlined in a page:
    `![](data:image/png;base64...)`.

    No fallback: there is no second HTML reader here, and a failure goes to
    `scan_report`, which names and counts it.
    """
    try:
        _markitdown()
    except ImportError:
        print(f"{RED}Error: markitdown not installed. Run: uv sync --extra documents{RESET}")
        return None
    kind = filepath.suffix.lstrip(".").upper()
    return _markdown_companion(
        filepath, kind, _markitdown_text(filepath, filepath.suffix.lower()))


def _extract_docx_native(filepath, fallback_reason: str | None = None):
    """Extract DOCX paragraphs AND tables to markdown, in document order.

    `Document.paragraphs` and `Document.tables` are separate collections. A
    reader that walks only the first returns nothing at all for a document
    whose content lives in tables (a rate card, a compliance matrix, a
    requirements sheet), and one that walks the second after the first reorders
    the document into two halves that were never adjacent. So the body's own
    child elements are walked instead, and a table is extracted where it
    actually sits.

    Capped at `MAX_EXTRACT_CHARS` and cut at a block boundary (a whole
    paragraph, or a whole table), never inside one.

    Nothing is caught here, for the reason given in `extract_pdf`.
    """
    try:
        import docx
        from docx.oxml.ns import qn
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError:
        print(f"{RED}Error: python-docx not installed. "
              f"Run: uv sync --extra documents{RESET}")
        return None

    document = docx.Document(str(filepath))
    children = list(document.element.body.iterchildren())

    body: list = []
    total_chars = 0
    paragraphs = 0
    tables = 0
    truncated_after = None

    for index, child in enumerate(children, 1):
        if total_chars >= MAX_EXTRACT_CHARS:
            truncated_after = index - 1
            break

        if child.tag == qn("w:p"):
            text = Paragraph(child, document).text.strip()
            if not text:
                continue
            body.append(text)
            body.append("")
            total_chars += len(text)
            paragraphs += 1
        elif child.tag == qn("w:tbl"):
            rows = _docx_table_rows(Table(child, document))
            if not rows:
                continue
            body.extend(rows)
            body.append("")
            total_chars += sum(len(r) for r in rows)
            tables += 1

    if not body:
        return _no_text_companion(
            filepath, "DOCX",
            "The document body holds no paragraph text and no tables.")

    note = None
    if truncated_after is not None:
        note = _truncation_note(
            MAX_EXTRACT_CHARS,
            f"block {truncated_after} of {len(children)}", "DOCX")

    lines = _extract_header(filepath, "DOCX")
    if fallback_reason:
        # See `_extract_pptx_native`: the provenance line names the table's
        # converter, this line says what actually read the file.
        lines += [f"> Read with python-docx: markitdown failed ({fallback_reason}).", ""]
    if note:
        lines.append(note)
        lines.append("")
    lines.append(f"Paragraphs: {paragraphs}")
    lines.append(f"Tables: {tables}")
    lines.append("")
    lines.extend(body)
    if note:
        lines.append(note)
        lines.append("")

    return "\n".join(lines)


class Converter(NamedTuple):
    """One row of the per-extension table.

    `name` is written into every companion and is what staleness compares, so
    replacing a converter regenerates exactly the companions the old one made
    and nothing else. `dist` is the installed distribution whose version is
    recorded beside the name.
    """
    name: str
    dist: str
    extract: Callable


# The converter for each extension, as HOS-148 measured it (see the module
# docstring). A row's NAME is what the provenance line records and what
# staleness compares, so a row that changes name regenerates exactly the
# companions the old one made. DOCX and PPTX name markitdown even when its
# failure hands a file to this script's own reader: the companion says so in
# its body, and the next change of source bytes tries markitdown again.
CONVERTERS: dict[str, Converter] = {
    ".docx": Converter("markitdown", "markitdown", extract_docx),
    ".htm": Converter("markitdown", "markitdown", extract_markitdown),
    ".html": Converter("markitdown", "markitdown", extract_markitdown),
    ".pdf": Converter("liteparse", "liteparse", extract_pdf),
    ".pptx": Converter("markitdown", "markitdown", extract_pptx),
    ".xlsx": Converter("openpyxl", "openpyxl", extract_xlsx),
}

EXTRACTABLE_SUFFIXES = tuple(CONVERTERS)


# ============================================================
# Provenance and staleness
# ============================================================

# The line every companion carries directly under its title since HOS-149.
_PROVENANCE = re.compile(
    r"^> Source SHA-256: (?P<sha>[0-9a-f]{64}); "
    r"converter: (?P<name>\S+) (?P<version>\S+)$", re.MULTILINE)

# How much of a companion is read to decide what it is. The title and the
# provenance line are the first three lines, so this is generous.
_HEAD_CHARS = 4096

MISSING = "missing"      # no file at the companion path
FRESH = "fresh"          # a companion of this source, current
STALE = "stale"          # a companion, but the source or the converter changed
LEGACY = "legacy"        # a companion written before provenance existed
OCCUPIED = "occupied"    # a file that is NOT this source's companion
FORCED = "forced"        # fresh, regenerated anyway under --force
UNCHANGED = "unchanged"  # source touched, bytes identical: fresh once re-dated

_WRITES = (MISSING, STALE, LEGACY, FORCED)

# The longest one document may take to convert, in seconds. The hourly
# memory-index refresh runs this before its build, so a parser that never
# returns would hold the index refresh with it. A document that hits the limit
# is a named failure and is retried on the next pass.
PER_FILE_SECONDS = 300


class _Overran(Exception):
    """A conversion ran past `PER_FILE_SECONDS`."""


def _convert_with_limit(conv: Converter, filepath: Path):
    """`conv.extract(filepath)`, abandoned after `PER_FILE_SECONDS`.

    SIGALRM, so the limit exists only on the main thread of a Unix process,
    which is where every caller of this script runs. Elsewhere it converts
    without a limit rather than refusing to convert.
    """
    import signal
    import threading
    if (not hasattr(signal, "SIGALRM")
            or threading.current_thread() is not threading.main_thread()):
        return conv.extract(filepath)

    def overran(signum, frame):
        raise _Overran(f"conversion exceeded {PER_FILE_SECONDS} s")

    previous = signal.signal(signal.SIGALRM, overran)
    signal.alarm(PER_FILE_SECONDS)
    try:
        return conv.extract(filepath)
    finally:
        signal.alarm(0)
        signal.signal(signal.SIGALRM, previous)


def source_sha256(path: Path) -> str:
    """The SHA-256 of a source document, streamed."""
    digest = hashlib.sha256()
    with open(path, "rb") as fh:
        for block in iter(lambda: fh.read(1 << 20), b""):
            digest.update(block)
    return digest.hexdigest()


def _stamp(content: str, sha: str, conv: Converter) -> str:
    """Insert the provenance line under the companion's title."""
    title, _, rest = content.partition("\n")
    line = (f"> Source SHA-256: {sha}; "
            f"converter: {conv.name} {dist_version(conv.dist)}")
    return f"{title}\n\n{line}\n{rest.lstrip(chr(10))}"


def companion_state(source: Path, companion: Path, conv: Converter) -> str:
    """What sits at `companion`, relative to `source`.

    A file is a companion of `source` only when its first line is the title
    every extractor writes, `# Extract: <source name>`. Anything else at that
    path is somebody's file and is never overwritten.

    The hash is compared only when the source is NEWER than its companion, so
    an hourly pass over two gigabytes of documents costs a stat per file rather
    than a read of every byte. The limit that buys, stated rather than left to
    be found: a source replaced by different bytes carrying an OLDER mtime (a
    restore with `cp -p`, an unzip) reads as fresh. `--force` regenerates it.
    """
    if not companion.exists():
        return MISSING
    with open(companion, encoding="utf-8", errors="replace") as fh:
        head = fh.read(_HEAD_CHARS)
    if head.partition("\n")[0] != f"# Extract: {source.name}":
        return OCCUPIED
    match = _PROVENANCE.search(head)
    if not match:
        return LEGACY
    if match["name"] != conv.name:
        return STALE
    if companion.stat().st_mtime >= source.stat().st_mtime:
        return FRESH
    return UNCHANGED if match["sha"] == source_sha256(source) else STALE


# ============================================================
# Which files are documents
# ============================================================

def _is_archive_name(name: str) -> bool:
    """True for any directory name that reads as an archive.

    Deliberately wide: `_archive`, `archive`, `linkedin-archive` and `v1-archive`
    all qualify. The operator's rule is that nothing is generated into anything
    that reads as an archive, and a companion written there is a new file inside
    a record meant to stay as it was. A name this over-matches costs a missing
    companion; one it under-matches costs a write into an archive.
    """
    return "archive" in name.lower()


# Walker policy of this script, beside the shared names of `excluded_dirs`.
_SKIP_DIRS = frozenset({"node_modules", "__pycache__"})

# Test fixtures are inputs to a test, not documents anyone recalls, and some are
# broken ON PURPOSE. MEASURED 2026-10-03 (HOS-194): the hourly walk reached a
# parked test suite in the data repository and tried a deliberately corrupt
# `.docx` fixture every hour, failing each time, and had already written three
# companions into fixture folders for the memory index to pick up. Matched
# case-insensitively, by the conventional names only. An explicit target still
# reaches a fixture: this prunes the WALK, not the operator.
_FIXTURE_DIRS = frozenset({"fixtures", "__fixtures__"})


def _skip_dir(name: str) -> bool:
    return (name.startswith(".") or _is_archive_name(name)
            or name in _SKIP_DIRS or name in SHARED_EXCLUDED_NAMES
            or name.lower() in _FIXTURE_DIRS)


def _is_document(path: Path) -> bool:
    """A file this script converts. `~$name.docx` is an Office lock file."""
    name = path.name
    return (path.suffix.lower() in CONVERTERS
            and not name.startswith((".", "~$")))


def _inside_archive(path: Path) -> bool:
    """Whether `path`, already known to be in the data repository, is in an archive."""
    rel = path.resolve().relative_to(get_data_root().resolve())
    return any(_is_archive_name(part) for part in rel.parts)


def document_roots() -> tuple[Path, ...]:
    """The top-level directories of the data repository a document may sit in.

    Every top-level directory in name order, so the C3 subject-tree roots are
    scanned without a list that has to grow with each move. MEASURED
    2026-10-13: 174 documents sat under subject-tree roots the family-root scan
    did not reach. Until HOS-240 (2026-10-13) the legacy family roots came
    first, in vocabulary order; they are now ordinary directories in the same
    name-ordered walk, so nothing here names them.

    Outside, as before: the runtime store (caches and scratch, under its
    registry name), and whatever `_skip_dir` prunes, which is dot directories,
    archives and the shared exclusions.
    """
    root = get_data_root()
    store = store_dirname(root)
    return tuple(p for p in sorted(root.iterdir())
                 if p.is_dir() and p.name != store and not _skip_dir(p.name))


def find_documents(scan_dirs) -> list:
    """Every document under `scan_dirs`, pruning what `_skip_dir` names."""
    found = []
    for top in scan_dirs:
        for dirpath, dirnames, filenames in os.walk(top):
            dirnames[:] = sorted(d for d in dirnames if not _skip_dir(d))
            for name in filenames:
                path = Path(dirpath) / name
                if _is_document(path):
                    found.append(path)
    return sorted(found)


def get_companion_path(filepath):
    """Get the -extract.md companion path for a binary file."""
    return filepath.with_name(filepath.stem + "-extract.md")


def _ambiguous_stems(paths) -> set:
    """(parent, stem) pairs claimed by more than one extractable file.

    The suffix is compared EXACTLY. Lowercased, `case.html` and `case.HTML`
    counted as one suffix, both claimed `case-extract.md`, and the second was
    skipped as "not its companion" on every run (reproduced in review
    2026-09-26).
    """
    by_stem: dict = {}
    for path in paths:
        by_stem.setdefault((path.parent, path.stem), set()).add(path.suffix)
    return {key for key, suffixes in by_stem.items() if len(suffixes) > 1}


def _companion_title(path: Path) -> str:
    """The first line of a companion, or "" when there is no file there."""
    if not path.is_file():
        return ""
    with open(path, encoding="utf-8", errors="replace") as fh:
        return fh.readline().rstrip("\n")


def _companion_for(filepath, ambiguous: set):
    """The companion path, with the source suffix added only when it must be.

    `<stem>-extract.md` is derived from the STEM, so `pitch.pptx` and
    `pitch.xlsx` in one folder claim the same companion. The deck was extracted
    first (it sorts first), and then the workbook was told "Skip pitch.xlsx
    (companion already exists)" - a companion that describes the deck. The
    workbook was never extracted and nothing said so. Under `--force` it was
    worse: both ran, the second overwrote the first, and the run reported two
    files extracted with one file of output.

    Only a colliding pair is renamed, and not the member the unsuffixed
    companion already names in its title: that file IS its companion, and
    renaming it would leave the old file behind as a duplicate of the new one.
    Admitting HTML made same-stem pairs common (a report rendered as both PDF
    and HTML), so without this the existing companion of every such pair would
    have been orphaned.
    """
    base = get_companion_path(filepath)
    if (filepath.parent, filepath.stem) not in ambiguous:
        return base
    if _companion_title(base) == f"# Extract: {filepath.name}":
        return base
    # The suffix as written, so case twins get two names. An uppercase suffix
    # that used to be lowercased here is renamed only when it shares its stem.
    suffix = filepath.suffix.lstrip(".")
    return base.with_name(f"{filepath.stem}-{suffix}-extract.md")


def orphaned_companions(ambiguous: set) -> list:
    """Unsuffixed companions left behind when a same-stem sibling ARRIVES.

    The rename above is scoped to currently-ambiguous stems on the stated
    assumption that "every companion on disk today belongs to a file with no
    same-stem sibling". That is true on the day it was written and stops being
    true the moment somebody drops `pitch.xlsx` beside an already-extracted
    `pitch.pptx`: both sources now resolve to suffixed companions, and the old
    `pitch-extract.md` -- which describes only the deck -- sits between them,
    unreferenced, for the next reader to take as the companion for either. The
    stem was printed in the ambiguity warning; the stale FILE never was.

    Reported, never deleted. It is extracted content and the operator may have
    edited it. A companion whose title names a same-stem document that exists is
    not an orphan: `_companion_for` keeps it as that document's companion.
    """
    orphans = []
    for parent, stem in sorted(ambiguous):
        stale = parent / f"{stem}-extract.md"
        if not stale.is_file():
            continue
        title = _companion_title(stale)
        named = title.removeprefix("# Extract: ")
        if (title != named and Path(named).stem == stem
                and (parent / named).is_file()):
            continue
        orphans.append(stale)
    return orphans


class ScanResult(NamedTuple):
    """What one scan produced, what it could not, and what it found on disk."""
    extracted: list
    failures: list
    states: Counter


def _targets(target):
    """(documents, the directory set ambiguity is judged over, a label), or None."""
    if target is None:
        dirs = document_roots()
        if not dirs:
            print(f"{RED}No document root exists under {get_data_root()}{RESET}")
            return None
        return (find_documents(dirs), None,
                f"the {len(dirs)} document roots of {get_data_root()}")
    path = Path(target)
    if not path.exists():
        print(f"{RED}Not found: {path}{RESET}")
        return None
    # A companion carries the document's text, so it may only ever be written
    # inside the data repository. An explicit target elsewhere, the engine
    # checkout above all, is refused: a companion written there is private
    # content in a public repository. The Telegram hook passes `--output`
    # paths through here, which is how this was reachable (review, 2026-09-26).
    if not path.resolve().is_relative_to(get_data_root().resolve()):
        print(f"{RED}Refused: {path} is outside the data repository "
              f"{get_data_root()}, and a companion is written only inside "
              f"it.{RESET}")
        return None
    if _inside_archive(path):
        print(f"{YELLOW}Refused: {path} is inside an archive, and no companion "
              f"is ever written into one.{RESET}")
        return None
    if path.is_file():
        if not _is_document(path):
            print(f"{YELLOW}Not a document this script converts: {path}{RESET}")
            return None
        # Ambiguity is a property of the folder, so the siblings decide the
        # companion's name even when one file is asked for.
        siblings = [p for p in path.parent.iterdir()
                    if p.is_file() and _is_document(p)]
        return [path], siblings, str(path)
    return find_documents((path,)), None, str(path)


def scan_report(target_dir=None, force=False, dry_run=False) -> ScanResult:
    """Write every missing or stale companion under the target, or count them.

    `target_dir` may be a directory or one document; None means every
    document root of the data repository (`document_roots`). `dry_run` reads
    and hashes but writes nothing.
    """
    found = _targets(target_dir)
    if found is None:
        return ScanResult([], [], Counter())
    documents, siblings, scan_label = found
    if not documents:
        print(f"{YELLOW}No documents found in {scan_label}{RESET}")
        return ScanResult([], [], Counter())

    ambiguous = _ambiguous_stems(siblings if siblings is not None else documents)
    if ambiguous:
        # The count every run; the suffixed companions themselves are named as
        # they are written, which is when the operator can first be surprised
        # by one. Listing every stem each hour buried the journal of the
        # refresh that runs this.
        print(f"{YELLOW}{len(ambiguous)} stem(s) are claimed by more than one "
              f"document; those companions carry the source suffix.{RESET}")
        orphans = orphaned_companions(ambiguous)
        if orphans:
            print(f"{YELLOW}{len(orphans)} unsuffixed companion(s) now describe "
                  f"only one of a pair. Read and delete or rename by hand:{RESET}")
            for stale in orphans:
                print(f"  {YELLOW}{stale}{RESET}")

    extracted = []
    failures: list[tuple[Path, str]] = []
    states: Counter = Counter()

    # Two documents whose companion names coincide (`p.html` suffixed to
    # `p-html-extract.md` beside a `p-html.htm`) would each see the other's
    # companion as occupied, forever, at exit 0. Both are FAILURES instead:
    # named, and counted in the exit code.
    claims: dict = {}
    for filepath in documents:
        claims.setdefault(_companion_for(filepath, ambiguous), []).append(filepath)
    for companion, claimants in claims.items():
        if len(claimants) > 1:
            names = ", ".join(c.name for c in claimants)
            for c in claimants:
                failures.append((c, f"companion name {companion.name} is "
                                    f"claimed by {names}; rename one"))
    collided = {c for _, cl in claims.items() if len(cl) > 1 for c in cl}
    documents = [d for d in documents if d not in collided]
    would_write: Counter = Counter()
    for filepath in documents:
        conv = CONVERTERS[filepath.suffix.lower()]
        companion = _companion_for(filepath, ambiguous)

        # Per file, not per batch, for the parse AND the write. A corrupt zip,
        # an encrypted PDF, an unreadable companion, a read-only folder or a
        # full disk is one file's problem; with nothing catching it here the
        # exception left the loop and every later document went unprocessed,
        # the run ending on a traceback rather than a report. MEASURED
        # 2026-09-02 with a chmod 555 subfolder: `PermissionError` escaped and a
        # good file sorting after it was never extracted.
        try:
            state = companion_state(filepath, companion, conv)
            if state == UNCHANGED:
                # Re-date the companion so the next pass is a stat again rather
                # than a full read of this source every hour. MEASURED
                # 2026-09-26 in review: a touched source was re-hashed on each
                # of three runs. The cost, stated: the memory index sees the
                # companion's mtime move and re-embeds it once.
                if not dry_run:
                    os.utime(companion)
                state = FRESH
            if force and state == FRESH:
                state = FORCED
            states[state] += 1
            if state == OCCUPIED:
                print(f"  {YELLOW}Skip{RESET}  {filepath.name}: {companion.name} "
                      f"exists and is not its companion")
                continue
            if state not in _WRITES:
                continue
            would_write[filepath.suffix.lower()] += 1
            if dry_run:
                continue
            print(f"  {BOLD}Extracting{RESET}  {filepath.name} ({state})...")
            sha = source_sha256(filepath)
            content = _convert_with_limit(conv, filepath)
            if not content:
                print(f"  {RED}Failed{RESET}  Could not extract {filepath.name}")
                failures.append((filepath, "extractor returned nothing"))
                continue
            atomic_write_text(companion, _stamp(content, sha, conv))
        except Exception as exc:  # noqa: BLE001 - one file, not the batch
            print(f"  {RED}Failed{RESET}  {filepath.name}: "
                  f"{type(exc).__name__}: {exc}")
            failures.append((filepath, f"{type(exc).__name__}: {exc}"))
            continue
        print(f"  {GREEN}Created{RESET}  {companion.name}")
        extracted.append((filepath, companion))

    print(f"\n{len(documents)} document(s) in {scan_label}: "
          + ", ".join(f"{states[s]} {s}" for s in
                      (FRESH, MISSING, STALE, LEGACY, FORCED, OCCUPIED)
                      if states[s]))
    if dry_run:
        print(f"Dry run, nothing written. Would write "
              f"{sum(would_write.values())}: "
              + ", ".join(f"{n} {ext}" for ext, n in sorted(would_write.items())))

    if failures:
        # Named, not swallowed: the operator can tell a run of 40 that produced
        # 40 from a run of 40 that produced 38.
        print(f"\n{YELLOW}{len(failures)} file(s) could not be extracted:{RESET}")
        for path, why in failures:
            print(f"  {YELLOW}{path.name}{RESET}: {why}")

    return ScanResult(extracted, failures, states)


def update_index(extracted_files):
    """Add newly extracted files to INDEX.md."""
    index = index_file()
    homes = datastore_homes()
    if not index.exists():
        print(f"{YELLOW}INDEX.md not found - skipping index update{RESET}")
        return

    try:
        content = index.read_text(encoding="utf-8")
    except (OSError, UnicodeDecodeError) as exc:
        # Same failure the comment below already names, one line earlier: this
        # runs LAST, after every file has been extracted, so raising here
        # reports failure over work that succeeded. `UnicodeDecodeError`
        # subclasses ValueError and this read had no handler at all.
        print(f"{YELLOW}INDEX.md unreadable ({exc}) - skipping index "
              f"update{RESET}")
        return
    today = datetime.now(get_default_tz()).strftime("%Y-%m-%d")

    new_rows = []
    for orig, _companion in extracted_files:
        # A target outside the DataStore is allowed on the CLI, and
        # `relative_to` raised ValueError for it — AFTER every file had already
        # been extracted, so the work was done and the index update died.
        # Rows are relative to the DataStore subtree the file sits in, which
        # is the pre-move spelling the move kept as a tail under each home.
        rel_path = next((orig.relative_to(h) for h in homes
                         if orig.is_relative_to(h)), None)
        if rel_path is None:
            print(f"{YELLOW}Skipping index row for {orig}: it is outside "
                  f"every DataStore home, which INDEX.md rows are relative "
                  f"to.{RESET}")
            continue
        domain = rel_path.parts[0] if rel_path.parts else "unknown"
        # Check if already in index
        if str(rel_path) in content:
            continue
        new_rows.append(f"| `{rel_path}` | {domain.title()} | {orig.stem} (auto-extracted) | {today} | *Review and update validates column* |")

    if not new_rows:
        print(f"{GREEN}INDEX.md already up to date{RESET}")
        return

    # Insert before the LAST HTML comment, not the first one anywhere in the
    # file. The comment above this said "before the closing comment", and the
    # code replaced the first `<!--` it found — so an INDEX.md opening with a
    # header note or a licence comment got data rows injected above it, before
    # the table they belong to.
    marker = content.rfind("<!--")
    block = "\n".join(new_rows)
    if marker != -1:
        content = content[:marker] + block + "\n\n" + content[marker:]
    else:
        content = content.rstrip("\n") + "\n" + block + "\n"

    # Rewrite the whole "Last updated" LINE, and carry its old value into a
    # single Previous line. Replacing only the literal prefix left the old date
    # dangling after `> Previous: `, and the next run did it again — nesting
    # one more stale Previous into the block every time.
    content = re.sub(
        r"^> Last updated:[^\n]*(?:\n> Previous:[^\n]*)?",
        lambda m: (f"> Last updated: {today}  \n> Previous: "
                   f"{m.group(0).split(':', 1)[1].split(chr(10))[0].strip()}"),
        content,
        count=1,
        flags=re.MULTILINE,
    )

    atomic_write_text(index, content)
    print(f"{GREEN}Added {len(new_rows)} entries to INDEX.md{RESET}")


def main():
    parser = argparse.ArgumentParser(description="Document companion extraction")
    parser.add_argument("target", nargs="?", default=None,
                        help="One directory or one document (default: every "
                             "top-level document directory of the data "
                             "repository)")
    parser.add_argument("--update-index", action="store_true",
                        help="Update INDEX.md with new DataStore entries")
    parser.add_argument("--force", action="store_true",
                        help="Regenerate every companion, fresh ones included")
    parser.add_argument("--dry-run", action="store_true",
                        help="Count what would be written; write nothing")
    args = parser.parse_args()

    print(f"\n{BOLD}Document companion extraction{RESET}\n")

    result = scan_report(args.target, force=args.force, dry_run=args.dry_run)

    if result.extracted and args.update_index:
        # INDEX.md rows are relative to a DataStore home, so only those rows
        # are offered; a companion elsewhere is indexed by the memory index.
        homes = datastore_homes()
        in_homes = [(o, c) for o, c in result.extracted
                    if any(o.is_relative_to(h) for h in homes)]
        update_index(in_homes)

    print(f"\n{BOLD}Done.{RESET} Extracted {len(result.extracted)} file(s).")

    # A run that could not read some of its input is not a successful run. The
    # failures were already named above; this is what lets a caller that is not
    # reading the output - a timer, a hook, a shell `&&` - tell the difference.
    if result.failures:
        print(f"{RED}{len(result.failures)} file(s) failed to extract.{RESET}")
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
